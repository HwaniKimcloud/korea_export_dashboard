from pathlib import Path
import pandas as pd
from openpyxl import load_workbook


# ============================================================
# 1. 기본 설정
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "industry_data"

INPUT_FILE = DATA_DIR / "2026 MTI-HSK 코드표_vFF_260507.xlsx"

if not INPUT_FILE.exists():
    raise FileNotFoundError(INPUT_FILE)


TARGETS = {
    "가전": [
        "가전",
        "가정용",
        "가정용전기",
        "가정용전자",
    ],
    "일반기계": [
        "일반기계",
        "산업기계",
        "제조장비",
        "에너지기계",
        "기계부품",
        "공작기계",
        "기계요소",
    ],
    "이차전지": [
        "이차전지",
        "2차전지",
        "배터리",
        "축전지",
        "리튬이온",
    ],
    "농수산식품": [
        "농수산식품",
        "농수산",
        "농산",
        "수산",
        "식품",
    ],
}


print()
print("=" * 110)
print("OFFICIAL MTI LAYOUT PROBE")
print("=" * 110)


# ============================================================
# 2. Workbook / Sheet 확인
# ============================================================

wb = load_workbook(
    INPUT_FILE,
    data_only=True
)

print()
print("[시트]")
for ws in wb.worksheets:
    print(
        f"- {ws.title}: "
        f"{ws.max_row:,}행 × {ws.max_column:,}열"
    )


# ============================================================
# 3. 병합셀 정보
# ============================================================

merge_rows = []

for ws in wb.worksheets:

    for merged_range in ws.merged_cells.ranges:

        merge_rows.append({
            "시트": ws.title,
            "병합범위": str(merged_range),
        })


merge_df = pd.DataFrame(merge_rows)

print()
print("=" * 110)
print("병합셀 개수")
print("=" * 110)

if merge_df.empty:
    print("병합셀 없음")
else:
    print(
        merge_df.groupby("시트")
        .size()
        .to_string()
    )


# ============================================================
# 4. 모든 시트에서 4개 산업 원문 검색
# ============================================================

matches = []


for ws in wb.worksheets:

    for row in ws.iter_rows():

        for cell in row:

            value = cell.value

            if value is None:
                continue

            text = str(value).strip()

            if not text:
                continue

            for industry, aliases in TARGETS.items():

                for alias in aliases:

                    if alias.lower() in text.lower():

                        matches.append({
                            "산업": industry,
                            "검색어": alias,
                            "시트": ws.title,
                            "행": cell.row,
                            "열": cell.column,
                            "셀주소": cell.coordinate,
                            "셀내용": text,
                        })


matches_df = pd.DataFrame(matches)


print()
print("=" * 110)
print("4개 산업 검색 결과")
print("=" * 110)

if matches_df.empty:

    print("검색 결과 없음")

else:

    summary = (
        matches_df
        .groupby(
            [
                "산업",
                "시트"
            ],
            as_index=False
        )
        .size()
    )

    print(
        summary.to_string(
            index=False
        )
    )


# ============================================================
# 5. 검색된 셀 주변 행 전체 출력
#
# 핵심:
# 해당 키워드가 발견된 행의
# 앞뒤 2행 + 모든 비어있지 않은 셀을 출력
# ============================================================

context_records = []


if not matches_df.empty:

    unique_hits = (
        matches_df[
            [
                "산업",
                "시트",
                "행"
            ]
        ]
        .drop_duplicates()
    )


    for _, hit in unique_hits.iterrows():

        industry = hit["산업"]
        sheet = hit["시트"]
        center_row = int(hit["행"])

        ws = wb[sheet]

        start_row = max(
            1,
            center_row - 2
        )

        end_row = min(
            ws.max_row,
            center_row + 2
        )


        for row_num in range(
            start_row,
            end_row + 1
        ):

            values = []

            for col_num in range(
                1,
                ws.max_column + 1
            ):

                value = ws.cell(
                    row=row_num,
                    column=col_num
                ).value

                if value is None:
                    continue

                text = str(value).strip()

                if text == "":
                    continue

                values.append(
                    f"C{col_num}={text}"
                )


            context_records.append({
                "산업": industry,
                "시트": sheet,
                "검색중심행": center_row,
                "현재행": row_num,
                "행내용": " | ".join(values),
            })


context_df = pd.DataFrame(
    context_records
)


# ============================================================
# 6. 산업별 대표 Context 출력
# ============================================================

print()
print("=" * 110)
print("산업별 원본 행 구조")
print("=" * 110)


for industry in TARGETS.keys():

    print()
    print("#" * 110)
    print(f"[{industry}]")
    print("#" * 110)

    if context_df.empty:

        print("없음")
        continue


    temp = (
        context_df[
            context_df["산업"]
            ==
            industry
        ]
        .drop_duplicates(
            subset=[
                "시트",
                "검색중심행",
                "현재행",
                "행내용"
            ]
        )
    )


    if temp.empty:

        print("없음")
        continue


    # 너무 길지 않게 우선 40행
    print(
        temp.head(40)
        .to_string(
            index=False
        )
    )


# ============================================================
# 7. 검색된 셀이 병합셀 안에 있는지도 확인
# ============================================================

merged_hit_records = []


if not matches_df.empty:

    for _, hit in matches_df.iterrows():

        ws = wb[hit["시트"]]

        coordinate = hit["셀주소"]

        merged_range_found = ""

        for merged_range in ws.merged_cells.ranges:

            if coordinate in merged_range:

                merged_range_found = str(
                    merged_range
                )

                break


        merged_hit_records.append({
            **hit.to_dict(),
            "병합범위": merged_range_found
        })


merged_hits_df = pd.DataFrame(
    merged_hit_records
)


# ============================================================
# 8. CSV 저장
# ============================================================

matches_df.to_csv(
    BASE_DIR / "official_mti_layout_hits.csv",
    index=False,
    encoding="utf-8-sig"
)

context_df.to_csv(
    BASE_DIR / "official_mti_layout_context.csv",
    index=False,
    encoding="utf-8-sig"
)

merged_hits_df.to_csv(
    BASE_DIR / "official_mti_layout_merged_hits.csv",
    index=False,
    encoding="utf-8-sig"
)

merge_df.to_csv(
    BASE_DIR / "official_mti_layout_all_merges.csv",
    index=False,
    encoding="utf-8-sig"
)


print()
print("=" * 110)
print("CSV 저장 완료")
print("=" * 110)

print("- official_mti_layout_hits.csv")
print("- official_mti_layout_context.csv")
print("- official_mti_layout_merged_hits.csv")
print("- official_mti_layout_all_merges.csv")


print()
print("=" * 110)
print("LAYOUT PROBE COMPLETE")
print("=" * 110)
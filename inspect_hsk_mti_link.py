from pathlib import Path
import pandas as pd
from openpyxl import load_workbook


# ============================================================
# 1. 경로 설정
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "industry_data"

INPUT_FILE = DATA_DIR / "2026 MTI-HSK 코드표_vFF_260507.xlsx"

if not INPUT_FILE.exists():
    raise FileNotFoundError(
        f"공식 코드표를 찾지 못했습니다:\n{INPUT_FILE}"
    )


print()
print("=" * 100)
print("HSK-MTI OFFICIAL LINK INSPECTOR")
print("=" * 100)


# ============================================================
# 2. Workbook 열기
# ============================================================

wb = load_workbook(
    INPUT_FILE,
    data_only=True,
    read_only=False
)


# ============================================================
# 3. HSK-MTI 연계표 시트 찾기
# ============================================================

target_sheet = None

for sheet_name in wb.sheetnames:

    normalized = (
        sheet_name
        .replace(" ", "")
        .replace("-", "")
        .upper()
    )

    if (
        "HSK" in normalized
        and "MTI" in normalized
        and "연계" in normalized
    ):
        target_sheet = sheet_name
        break


if target_sheet is None:
    raise ValueError(
        "HSK-MTI 연계표 시트를 찾지 못했습니다."
    )


ws = wb[target_sheet]


print()
print(f"사용 시트 : {target_sheet}")
print(f"전체 행   : {ws.max_row:,}")
print(f"전체 열   : {ws.max_column:,}")


# ============================================================
# 4. 맨 위 30행 원본 그대로 출력
# ============================================================

print()
print("=" * 100)
print("STEP 1. 상단 30행 원본 구조")
print("=" * 100)


for row_num in range(
    1,
    min(ws.max_row, 30) + 1
):

    values = []

    for col_num in range(
        1,
        ws.max_column + 1
    ):

        value = ws.cell(
            row=row_num,
            column=col_num
        ).value

        if value is None:
            continue

        text = str(value).strip()

        if text == "":
            continue

        values.append(
            f"C{col_num}={text}"
        )


    if values:
        print(
            f"ROW {row_num:>3} | "
            + " | ".join(values)
        )


# ============================================================
# 5. 병합 셀 확인
# ============================================================

print()
print("=" * 100)
print("STEP 2. 병합 셀")
print("=" * 100)

merged_ranges = list(
    ws.merged_cells.ranges
)

print(
    f"병합 셀 범위 수: {len(merged_ranges):,}"
)

for merged_range in merged_ranges[:50]:
    print(
        f"- {merged_range}"
    )


# ============================================================
# 6. HSK / MTI처럼 보이는 데이터 행 자동 탐색
# ============================================================

print()
print("=" * 100)
print("STEP 3. HSK-MTI 데이터 행 후보")
print("=" * 100)


candidate_rows = []


for row_num in range(
    1,
    ws.max_row + 1
):

    row_values = []

    for col_num in range(
        1,
        ws.max_column + 1
    ):

        value = ws.cell(
            row=row_num,
            column=col_num
        ).value

        if value is None:
            row_values.append("")
        else:
            row_values.append(
                str(value).strip()
            )


    # 숫자 코드 후보
    numeric_values = []

    for col_num, value in enumerate(
        row_values,
        start=1
    ):

        cleaned = (
            value
            .replace(",", "")
            .replace(".0", "")
            .strip()
        )

        if cleaned.isdigit():

            numeric_values.append(
                (
                    col_num,
                    cleaned
                )
            )


    # HSK는 보통 10자리,
    # MTI는 1~6자리 숫자 코드
    has_hsk = any(
        len(code) == 10
        for _, code in numeric_values
    )

    has_mti = any(
        1 <= len(code) <= 6
        for _, code in numeric_values
    )


    if has_hsk and has_mti:

        candidate_rows.append(
            (
                row_num,
                row_values
            )
        )


print(
    f"HSK+MTI 형태 후보 행: "
    f"{len(candidate_rows):,}"
)


print()
print("처음 20개 후보:")
print("-" * 100)


for row_num, row_values in candidate_rows[:20]:

    values = []

    for col_num, value in enumerate(
        row_values,
        start=1
    ):

        if value != "":
            values.append(
                f"C{col_num}={value}"
            )


    print(
        f"ROW {row_num:>5} | "
        + " | ".join(values)
    )


# ============================================================
# 7. 컬럼별 코드 형태 분석
# ============================================================

print()
print("=" * 100)
print("STEP 4. 컬럼별 코드 길이 분석")
print("=" * 100)


column_stats = []


for col_num in range(
    1,
    ws.max_column + 1
):

    length_counts = {}

    non_empty = 0

    for row_num in range(
        1,
        ws.max_row + 1
    ):

        value = ws.cell(
            row=row_num,
            column=col_num
        ).value

        if value is None:
            continue

        text = (
            str(value)
            .strip()
            .replace(",", "")
            .replace(".0", "")
        )

        if text == "":
            continue

        non_empty += 1

        if text.isdigit():

            code_length = len(text)

            length_counts[
                code_length
            ] = (
                length_counts.get(
                    code_length,
                    0
                )
                + 1
            )


    column_stats.append({
        "열": col_num,
        "비어있지않은셀": non_empty,
        "숫자길이분포": str(
            length_counts
        )
    })


stats_df = pd.DataFrame(
    column_stats
)

print(
    stats_df.to_string(
        index=False
    )
)


# ============================================================
# 8. CSV로 원본 시트 저장
# ============================================================

print()
print("=" * 100)
print("STEP 5. 원본 연계표 CSV 저장")
print("=" * 100)


raw_df = pd.read_excel(
    INPUT_FILE,
    sheet_name=target_sheet,
    header=None,
    dtype=str
)


OUTPUT_FILE = (
    BASE_DIR
    / "official_hsk_mti_link_raw.csv"
)


raw_df.to_csv(
    OUTPUT_FILE,
    index=False,
    encoding="utf-8-sig"
)


print(
    f"저장 완료: {OUTPUT_FILE.name}"
)


# ============================================================
# 9. 종료
# ============================================================

print()
print("=" * 100)
print("INSPECTION COMPLETE")
print("=" * 100)